
from io import BytesIO
from docx import Document

def create_docx_report(name_a, name_b, result):
    document = Document()

    document.add_heading(
        "AI Document Comparison Assistant", level=0
    )

    document.add_paragraph(f"Document A: {name_a}")
    document.add_paragraph(f"Document B: {name_b}")

    document.add_paragraph(
        "Informational comparison only. "
        "Verify important details against the original PDFs."
    )

    sections = [
        ("Document A Summary", "summary_a"),
        ("Document B Summary", "summary_b"),
        ("Similarities", "similarities"),
        ("Differences", "differences"),
        ("Key Details Comparison", "key_details"),
        ("Overall Conclusion", "conclusion"),
        ("Limitations", "limitations")
    ]

    for heading, key in sections:
        document.add_heading(heading, level=1)

        content = result.get(
            key, "Not found in the provided text."
        )

        for line in content.splitlines():
            if line.strip():
                document.add_paragraph(line.strip())

    output = BytesIO()
    document.save(output)

    return output.getvalue()
